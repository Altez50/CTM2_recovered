from __future__ import annotations

import csv
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.config import settings
from api.database import new_uuid, utcnow_iso
from api.tests.conftest import login_headers
from api import text_import_export


def _create_estimate(client: TestClient, headers: dict[str, str]) -> int:
    response = client.post(
        "/api/documents/estimates",
        headers=headers,
        json={"number": "P0-EST-1", "date": "2026-09-02", "lines": []},
    )
    assert response.status_code == 201, response.text
    return int(response.json()["data"]["id"])


def test_login_session_and_invalid_password(client: TestClient) -> None:
    assert client.post(
        "/api/auth/login", json={"username": "admin", "password": "wrong"}
    ).status_code == 401
    headers = login_headers(client)
    me = client.get("/api/auth/me", headers=headers)
    assert me.status_code == 200
    assert me.json()["role"] == "admin"
    assert client.post("/api/auth/logout", headers=headers).status_code == 200
    assert client.get("/api/auth/me", headers=headers).status_code == 401


def test_import_preview_apply_lifecycle_and_role_guards(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    estimate_id = _create_estimate(client, admin_headers)
    payload = json.dumps(
        [{"work_code": "P0-W1", "work_name": "P0 work", "unit": "m2", "quantity": 2, "price": 12.5, "sum": 25}],
        ensure_ascii=False,
    )

    preview = client.post(
        f"/api/documents/estimates/{estimate_id}/import-text?format=json&preview=true",
        content=payload.encode(),
        headers=admin_headers,
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["data"] == {
        "imported": 1,
        "warnings": [],
        "format": "json",
        "preview": True,
    }
    assert client.get(
        f"/api/documents/estimates/{estimate_id}", headers=admin_headers
    ).json()["data"]["lines"] == []

    foreman_headers = login_headers(client, "foreman", "foreman")
    forbidden = client.post(
        f"/api/documents/estimates/{estimate_id}/import-text?format=json",
        content=payload.encode(),
        headers=foreman_headers,
    )
    assert forbidden.status_code == 403

    applied = client.post(
        f"/api/documents/estimates/{estimate_id}/import-text?format=json",
        content=payload.encode(),
        headers=admin_headers,
    )
    assert applied.status_code == 200, applied.text
    estimate = client.get(
        f"/api/documents/estimates/{estimate_id}", headers=admin_headers
    ).json()["data"]
    assert len(estimate["lines"]) == 1
    assert estimate["lines"][0]["work_name"] == "P0 work"

    assert client.post(
        f"/api/documents/estimates/{estimate_id}/post", headers=admin_headers
    ).status_code == 200
    locked = client.post(
        f"/api/documents/estimates/{estimate_id}/import-text?format=json",
        content=payload.encode(),
        headers=admin_headers,
    )
    assert locked.status_code == 409
    edit = client.put(
        f"/api/documents/estimates/{estimate_id}",
        headers=admin_headers,
        json={"number": "changed", "date": "2026-09-02", "lines": []},
    )
    assert edit.status_code == 400
    assert client.post(
        f"/api/documents/estimates/{estimate_id}/unpost", headers=admin_headers
    ).status_code == 200


def test_sync_token_gate(client: TestClient) -> None:
    original = settings.sync_token
    settings.sync_token = "test-sync-token-with-enough-entropy"
    try:
        assert client.get("/api/sync/status").status_code == 401
        assert client.get(
            "/api/sync/status", headers={"X-Sync-Token": settings.sync_token}
        ).status_code == 200
    finally:
        settings.sync_token = original


def test_import_request_is_rejected_before_exceeding_memory_limit(
    client: TestClient,
    admin_headers: dict[str, str],
    monkeypatch,
) -> None:
    estimate_id = _create_estimate(client, admin_headers)
    monkeypatch.setattr(text_import_export, "MAX_IMPORT_BYTES", 8)
    response = client.post(
        f"/api/documents/estimates/{estimate_id}/import-text?format=json",
        content=b"[123456789]",
        headers=admin_headers,
    )
    assert response.status_code == 413


def test_daily_report_export_uses_schema_fields_for_tsv_and_xlsx(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    report = client.get("/api/documents/daily-reports/1", headers=admin_headers)
    assert report.status_code == 200, report.text
    assert len(report.json()["data"]["lines"]) == 2

    tsv = client.get(
        "/api/cost/daily-reports/1/export-text?format=tsv",
        headers=admin_headers,
    )
    assert tsv.status_code == 200, tsv.text
    tsv_rows = list(csv.reader(tsv.text.splitlines(), delimiter="\t"))
    assert len(tsv_rows) == 3
    assert tsv_rows[1][0:2] == ["W-001", "Brickwork"]
    assert tsv_rows[1][3] == "120"

    xlsx = client.get(
        "/api/cost/daily-reports/1/export-text?format=xlsx",
        headers=admin_headers,
    )
    assert xlsx.status_code == 200, xlsx.text
    sheet = load_workbook(io.BytesIO(xlsx.content), read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[1][0:2] == ("W-001", "Brickwork")
    assert rows[1][3] == 120


def test_import_rolls_back_lines_and_created_references_on_error(connection, monkeypatch) -> None:
    now = utcnow_iso()
    connection.execute(
        """INSERT INTO estimates (
               number, date, total_sum, total_labor, is_posted, marked_for_deletion,
               created_at, modified_at, uuid, updated_at, is_deleted
           ) VALUES ('P0-ROLLBACK', '2026-09-02', 1, 0, 0, 0, ?, ?, ?, ?, 0)""",
        (now, now, new_uuid(), now),
    )
    estimate_id = connection.execute("SELECT last_insert_rowid()").fetchone()[0]
    connection.execute(
        """INSERT INTO estimate_lines (
               estimate_id, line_number, quantity, price, sum, planned_labor,
               is_group, is_collapsed, marked_for_deletion, created_at,
               modified_at, uuid, updated_at, is_deleted
           ) VALUES (?, 1, 1, 1, 1, 0, 0, 0, 0, ?, ?, ?, ?, 0)""",
        (estimate_id, now, now, new_uuid(), now),
    )
    connection.commit()

    original = text_import_export._row_to_line
    calls = 0

    def fail_on_second(*args, **kwargs):
        nonlocal calls
        calls += 1
        result = original(*args, **kwargs)
        if calls == 2:
            raise RuntimeError("synthetic row failure")
        return result

    monkeypatch.setattr(text_import_export, "_row_to_line", fail_on_second)
    payload = json.dumps([
        {"work_code": "ROLLBACK-1", "work_name": "Rollback one", "quantity": 1},
        {"work_code": "ROLLBACK-2", "work_name": "Rollback two", "quantity": 1},
    ])
    with pytest.raises(RuntimeError, match="synthetic row failure"):
        text_import_export.import_document_lines(
            connection, "estimate", estimate_id, fmt="json", data=payload
        )

    assert connection.execute(
        "SELECT COUNT(*) FROM estimate_lines WHERE estimate_id = ?", (estimate_id,)
    ).fetchone()[0] == 1
    assert connection.execute(
        "SELECT COUNT(*) FROM works WHERE code LIKE 'ROLLBACK-%'"
    ).fetchone()[0] == 0
